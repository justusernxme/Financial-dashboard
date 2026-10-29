"""Unit tests for the Financial Statements Extractor and Excel Generator.

Tests XBRL facts extraction, DataFrame alignment across 10 years, Excel
workbook creation with 3 tabs, and workbook verification.
"""

import tempfile
import unittest
from pathlib import Path

import openpyxl

from financial_statements_extractor import (
    BALANCE_SHEET_TEMPLATE,
    CASH_FLOW_TEMPLATE,
    INCOME_STATEMENT_TEMPLATE,
    build_statement_dataframe,
    determine_target_years,
    export_financial_statements_to_excel,
    extract_concept_series,
    generate_mock_company_facts,
    run_financial_analysis_pipeline,
    verify_financial_excel_workbook,
)


class TestFinancialStatementsExtractor(unittest.TestCase):
    """Test data alignment, Excel styling, and verification."""

    def setUp(self) -> None:
        """Initialize mock financial data for 10 fiscal years."""
        self.temp_dir = tempfile.TemporaryDirectory()
        self.base_path = Path(self.temp_dir.name)
        self.ticker = "AAPL"
        self.years = [2016 + i for i in range(10)]
        self.mock_data = generate_mock_company_facts(self.ticker, self.years)
        self.facts_dict = self.mock_data["facts"]["us-gaap"]

    def tearDown(self) -> None:
        """Clean up temporary test directory."""
        self.temp_dir.cleanup()

    def test_determine_target_years(self) -> None:
        """Verify calculation of the 10 target fiscal years."""
        extracted_years = determine_target_years(self.facts_dict, num_years=10)
        self.assertEqual(len(extracted_years), 10)
        self.assertEqual(extracted_years[0], 2016)
        self.assertEqual(extracted_years[-1], 2025)

    def test_extract_concept_series(self) -> None:
        """Verify extraction of annual values for a specific XBRL concept."""
        tags = [
            "RevenueFromContractWithCustomerExcludingAssessedTax",
            "SalesRevenueNet",
        ]
        series = extract_concept_series(self.facts_dict, tags, self.years)
        self.assertEqual(len(series), 10)
        for yr in self.years:
            self.assertIn(yr, series)
            self.assertGreater(series[yr], 0)

    def test_build_statement_dataframes(self) -> None:
        """Verify creation of aligned DataFrames for all 3 statements."""
        income_df = build_statement_dataframe(
            template=INCOME_STATEMENT_TEMPLATE,
            us_gaap_facts=self.facts_dict,
            target_years=self.years,
        )
        self.assertIn("Line Item", income_df.columns)
        for yr in self.years:
            self.assertIn(str(yr), income_df.columns)

        balance_df = build_statement_dataframe(
            template=BALANCE_SHEET_TEMPLATE,
            us_gaap_facts=self.facts_dict,
            target_years=self.years,
        )
        self.assertEqual(len(balance_df.columns), 11)

        cash_df = build_statement_dataframe(
            template=CASH_FLOW_TEMPLATE,
            us_gaap_facts=self.facts_dict,
            target_years=self.years,
        )
        self.assertEqual(len(cash_df.columns), 11)

    def test_excel_export_and_verification(self) -> None:
        """Verify creation of 3-tab workbook and verification check."""
        income_df = build_statement_dataframe(
            template=INCOME_STATEMENT_TEMPLATE,
            us_gaap_facts=self.facts_dict,
            target_years=self.years,
        )
        balance_df = build_statement_dataframe(
            template=BALANCE_SHEET_TEMPLATE,
            us_gaap_facts=self.facts_dict,
            target_years=self.years,
        )
        cash_df = build_statement_dataframe(
            template=CASH_FLOW_TEMPLATE,
            us_gaap_facts=self.facts_dict,
            target_years=self.years,
        )

        excel_path = self.base_path / f"{self.ticker}_Financial_Analysis.xlsx"
        saved_path = export_financial_statements_to_excel(
            ticker=self.ticker,
            income_df=income_df,
            balance_df=balance_df,
            cash_df=cash_df,
            target_years=self.years,
            output_path=excel_path,
        )
        self.assertTrue(saved_path.is_file())

        # Load workbook with openpyxl and inspect sheets
        wb = openpyxl.load_workbook(saved_path, data_only=True)
        self.assertIn("Income Statement", wb.sheetnames)
        self.assertIn("Balance Sheet", wb.sheetnames)
        self.assertIn("Cash Flows", wb.sheetnames)

        # Check year headers on row 3
        ws_income = wb["Income Statement"]
        header_years = [
            ws_income.cell(row=3, column=col).value for col in range(2, 12)
        ]
        self.assertEqual(header_years, self.years)

        # Run verification engine
        is_valid = verify_financial_excel_workbook(
            excel_path=saved_path,
            expected_years=self.years,
        )
        self.assertTrue(is_valid)

    def test_run_financial_analysis_pipeline(self) -> None:
        """Verify the full end-to-end pipeline with mock data fallback."""
        excel_path = self.base_path / "AAPL_Analysis_Test.xlsx"
        success = run_financial_analysis_pipeline(
            ticker=self.ticker,
            output_path=excel_path,
            num_years=10,
            use_mock_on_network_failure=True,
        )
        self.assertTrue(success)
        self.assertTrue(excel_path.is_file())


if __name__ == "__main__":
    unittest.main()
