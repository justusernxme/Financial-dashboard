"""Unit tests for the Financial Ratios and Executive Summary Module.

Tests formula calculations for Liquidity, Solvency, Efficiency,
and Profitability ratios, Pearson correlation, sheet generation,
dual-axis chart construction, tab reordering, and workbook verification.
"""

import tempfile
from pathlib import Path
from typing import Dict, List, Optional
import unittest

import openpyxl

from financial_ratios_analyzer import (
    RATIO_CATALOG,
    ExecutiveSummaryRow,
    RatioDefinition,
    calculate_pearson_correlation,
    compute_ratio_series,
    generate_mock_financial_data,
    normalize_concept_label,
    reorder_sheets,
    run_financial_ratios_pipeline,
    safe_divide,
    verify_ratios_and_summary_tabs,
    write_executive_summary_sheet,
    write_financial_ratios_sheet,
)


class TestFinancialRatiosAnalyzer(unittest.TestCase):
    """Test suite for financial ratios engine and executive summary."""

    def setUp(self) -> None:
        """Initialize temporary workspace and test parameters."""
        self.temp_dir = tempfile.TemporaryDirectory()
        self.base_path = Path(self.temp_dir.name)
        self.ticker = "AAPL"
        self.years = [2021, 2022, 2023, 2024]
        self.mock_data = generate_mock_financial_data(self.ticker, self.years)

    def tearDown(self) -> None:
        """Clean up temporary test directory."""
        self.temp_dir.cleanup()

    def test_normalize_concept_label(self) -> None:
        """Verify normalization strips symbols, spacing, and casing."""
        raw = "   Total Net Sales / Revenue (USD) \n"
        norm = normalize_concept_label(raw)
        self.assertEqual(norm, "total net sales revenue usd")

    def test_safe_divide(self) -> None:
        """Verify division safely handles zeros, Nones, and floats."""
        self.assertAlmostEqual(safe_divide(10.0, 2.0), 5.0)
        self.assertIsNone(safe_divide(10.0, 0.0))
        self.assertIsNone(safe_divide(None, 2.0))
        self.assertIsNone(safe_divide(10.0, None))
        self.assertIsNone(safe_divide(0.0, 0.0))

    def test_pearson_correlation(self) -> None:
        """Verify Pearson correlation calculation between paired series."""
        # Perfect positive correlation
        r_pos = calculate_pearson_correlation([1.0, 2.0, 3.0], [2.0, 4.0, 6.0])
        self.assertIsNotNone(r_pos)
        self.assertAlmostEqual(r_pos, 1.0)

        # Perfect negative correlation
        r_neg = calculate_pearson_correlation([1.0, 2.0, 3.0], [6.0, 4.0, 2.0])
        self.assertIsNotNone(r_neg)
        self.assertAlmostEqual(r_neg, -1.0)

        # Insufficient data
        self.assertIsNone(calculate_pearson_correlation([1.0], [2.0]))

    def test_compute_all_ratios(self) -> None:
        """Verify computation across all four financial ratio categories."""
        for r_def in RATIO_CATALOG:
            res = compute_ratio_series(
                ratio_def=r_def,
                financial_data=self.mock_data,
                years=self.years,
            )
            # Must compute a value for every test year
            for yr in self.years:
                self.assertIn(yr, res)
                val = res[yr]
                self.assertIsNotNone(
                    val,
                    f"Ratio '{r_def.name}' returned None for year {yr}",
                )
                self.assertGreater(val, 0.0)

    def test_write_financial_ratios_sheet(self) -> None:
        """Verify 'Financial Ratios' sheet formatting, columns, and rows."""
        wb = openpyxl.Workbook()
        ratio_results: Dict[str, Dict[int, Optional[float]]] = {}
        for r_def in RATIO_CATALOG:
            ratio_results[r_def.name] = compute_ratio_series(
                ratio_def=r_def,
                financial_data=self.mock_data,
                years=self.years,
            )

        write_financial_ratios_sheet(
            wb=wb,
            ticker=self.ticker,
            years=self.years,
            ratio_results=ratio_results,
        )
        self.assertIn("Financial Ratios", wb.sheetnames)

        ws = wb["Financial Ratios"]
        self.assertIn("10-YEAR HISTORICAL FINANCIAL RATIOS", ws["A1"].value)
        self.assertEqual(
            ws.cell(row=3, column=1).value, "Ratio Category / Metric"
        )
        self.assertEqual(ws.cell(row=3, column=3).value, 2021)

        # Check that categories exist
        category_found = any(
            ws.cell(row=r, column=1).value == "LIQUIDITY RATIOS"
            for r in range(4, ws.max_row + 1)
        )
        self.assertTrue(category_found)

    def test_write_executive_summary_sheet_with_chart(self) -> None:
        """Verify 'Executive Summary' tab generation with dual-axis chart."""
        wb = openpyxl.Workbook()
        summary_rows = [
            ExecutiveSummaryRow(
                fiscal_year=2021,
                stock_price=146.92,
                yoy_return=30.85,
                pn_ratio=1.75,
                sentiment_tone="Positive",
                revenue=365000000000.0,
                net_income=94000000000.0,
                operating_margin=0.29,
                roe=1.45,
                current_ratio=1.07,
                debt_to_equity=1.75,
            ),
            ExecutiveSummaryRow(
                fiscal_year=2022,
                stock_price=150.43,
                yoy_return=2.39,
                pn_ratio=1.65,
                sentiment_tone="Positive",
                revenue=394000000000.0,
                net_income=99000000000.0,
                operating_margin=0.30,
                roe=1.96,
                current_ratio=0.88,
                debt_to_equity=2.36,
            ),
        ]

        write_executive_summary_sheet(
            wb=wb,
            ticker=self.ticker,
            summary_rows=summary_rows,
        )
        self.assertIn("Executive Summary", wb.sheetnames)

        ws = wb["Executive Summary"]
        # Header banner
        self.assertIn("EXECUTIVE SUMMARY", ws["A1"].value)
        # Check chart presence
        self.assertEqual(len(ws._charts), 1)
        chart = ws._charts[0]
        chart_title_text = chart.title.tx.rich.p[0].r[0].t
        self.assertIn("Stock Price vs. MD&A Sentiment", chart_title_text)

    def test_reorder_sheets(self) -> None:
        """Verify sheets are reordered into the executive logical sequence."""
        wb = openpyxl.Workbook()
        wb.create_sheet("Stock Price")
        wb.create_sheet("Balance Sheet")
        wb.create_sheet("Executive Summary")
        wb.create_sheet("Financial Ratios")
        wb.create_sheet("Income Statement")

        reorder_sheets(wb)
        expected_prefix = [
            "Executive Summary",
            "Financial Ratios",
            "Income Statement",
            "Balance Sheet",
            "Stock Price",
        ]
        self.assertEqual(wb.sheetnames[:5], expected_prefix)

    def test_pipeline_execution(self) -> None:
        """Verify full end-to-end pipeline execution in test environment."""
        wb_path = self.base_path / "AAPL_Financial_Analysis.xlsx"
        success = run_financial_ratios_pipeline(
            ticker="AAPL",
            workbook_path=wb_path,
            data_dir=self.base_path,
            num_years=5,
        )
        self.assertTrue(success)
        self.assertTrue(wb_path.is_file())

        verified = verify_ratios_and_summary_tabs(wb_path, expected_years=5)
        self.assertTrue(verified)


if __name__ == "__main__":
    unittest.main()
