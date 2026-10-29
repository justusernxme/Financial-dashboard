"""Unit tests for the Stock Price Extractor Module.

Tests balance sheet date identification, closest prior trading day logic,
stock price retrieval via yfinance (and mock fallbacks), Excel tab generation,
styling, and verification.
"""

import datetime
import json
from pathlib import Path
import tempfile
from typing import Any
import unittest
from unittest.mock import MagicMock, patch

import openpyxl
import pandas as pd

from stock_price_extractor import (
    StockPriceRecord,
    fetch_price_for_date,
    generate_mock_stock_prices,
    get_balance_sheet_dates,
    retrieve_balance_sheet_stock_prices,
    run_stock_price_pipeline,
    verify_stock_price_tab,
    write_stock_prices_to_excel,
)


class TestStockPriceExtractor(unittest.TestCase):
    """Test suite for stock price extraction and Excel export."""

    def setUp(self) -> None:
        """Initialize temporary workspace and test parameters."""
        self.temp_dir = tempfile.TemporaryDirectory()
        self.base_path = Path(self.temp_dir.name)
        self.ticker = "AAPL"
        self.years = [2021, 2022, 2023, 2024]

    def tearDown(self) -> None:
        """Clean up temporary resources."""
        self.temp_dir.cleanup()

    def test_get_balance_sheet_dates_fiscal_conventions(self) -> None:
        """Verify fiscal year-end date rules for standard tickers."""
        # Apple: last Saturday of September
        dates_aapl = get_balance_sheet_dates("AAPL", [2023, 2024])
        # 2023-09-30 was a Saturday
        self.assertEqual(dates_aapl[2023], datetime.date(2023, 9, 30))
        # 2024-09-28 was the last Saturday of September 2024
        self.assertEqual(dates_aapl[2024], datetime.date(2024, 9, 28))

        # Microsoft: June 30
        dates_msft = get_balance_sheet_dates("MSFT", [2024])
        self.assertEqual(dates_msft[2024], datetime.date(2024, 6, 30))

        # Generic calendar year: December 31
        dates_gen = get_balance_sheet_dates("XYZ", [2024])
        self.assertEqual(dates_gen[2024], datetime.date(2024, 12, 31))

    def test_get_balance_sheet_dates_from_cache(self) -> None:
        """Verify date extraction from cached SEC company facts JSON."""
        cache_dir = self.base_path / "TEST" / "cache"
        cache_dir.mkdir(parents=True, exist_ok=True)
        facts_file = cache_dir / "CIK0001234567_facts.json"

        mock_facts = {
            "facts": {
                "us-gaap": {
                    "Assets": {
                        "units": {
                            "USD": [
                                {
                                    "end": "2023-09-29",
                                    "fy": 2023,
                                    "form": "10-K",
                                }
                            ]
                        }
                    }
                }
            }
        }
        facts_file.write_text(json.dumps(mock_facts), encoding="utf-8")

        dates = get_balance_sheet_dates(
            ticker="TEST",
            years=[2023],
            data_dir=self.base_path,
        )
        self.assertEqual(dates[2023], datetime.date(2023, 9, 29))

    def test_fetch_price_closest_prior_trading_day(self) -> None:
        """Verify picking the latest trading day on or before target."""
        # Simulate target date on Sunday, Sept 29, 2024
        target_sunday = datetime.date(2024, 9, 29)

        # Mock yfinance history DataFrame
        dates = [
            pd.Timestamp("2024-09-26"),
            pd.Timestamp("2024-09-27"),  # Friday
            pd.Timestamp("2024-09-30"),  # Monday
        ]
        mock_df = pd.DataFrame(
            {
                "Close": [225.0, 227.5, 230.0],
                "Adj Close": [224.5, 227.0, 229.5],
                "Volume": [45000000, 50000000, 48000000],
            },
            index=pd.DatetimeIndex(dates),
        )

        mock_ticker = MagicMock()
        mock_ticker.history.return_value = mock_df

        trade_dt, close, adj_close, vol = fetch_price_for_date(
            ticker="AAPL",
            target_date=target_sunday,
            ticker_obj=mock_ticker,
        )

        # Should choose Friday 2024-09-27 as the closest prior trading day
        self.assertEqual(trade_dt, datetime.date(2024, 9, 27))
        self.assertAlmostEqual(close, 227.5)
        self.assertAlmostEqual(adj_close, 227.0)
        self.assertEqual(vol, 50000000)

    def test_generate_mock_stock_prices(self) -> None:
        """Verify synthetic price generation and YoY return calculation."""
        bs_dates = {
            2021: datetime.date(2021, 9, 25),
            2022: datetime.date(2022, 9, 24),
            2023: datetime.date(2023, 9, 30),
        }
        records = generate_mock_stock_prices("AAPL", bs_dates)
        self.assertEqual(len(records), 3)

        # First record has no prior year return
        self.assertIsNone(records[0].yoy_return)
        # Subsequent records have valid returns
        self.assertIsNotNone(records[1].yoy_return)
        self.assertIsNotNone(records[2].yoy_return)

        # Trading dates must be weekdays (Monday=0 ... Friday=4)
        for rec in records:
            self.assertLessEqual(rec.trading_date.weekday(), 4)
            self.assertLessEqual(rec.trading_date, rec.balance_sheet_date)

    def test_write_stock_prices_to_excel(self) -> None:
        """Verify creation of 'Stock Price' tab with corporate formatting."""
        wb_path = self.base_path / "AAPL_Financial_Analysis.xlsx"

        # Create workbook with existing statement tab
        wb = openpyxl.Workbook()
        ws_initial = wb.active
        ws_initial.title = "Income Statement"
        ws_initial["A1"] = "Revenue"
        wb.save(wb_path)

        # Generate sample records
        records = [
            StockPriceRecord(
                fiscal_year=2022,
                balance_sheet_date=datetime.date(2022, 9, 24),
                trading_date=datetime.date(2022, 9, 23),
                close_price=150.43,
                adj_close=149.20,
                volume=85000000,
                yoy_change=None,
                yoy_return=None,
            ),
            StockPriceRecord(
                fiscal_year=2023,
                balance_sheet_date=datetime.date(2023, 9, 30),
                trading_date=datetime.date(2023, 9, 29),
                close_price=171.21,
                adj_close=170.80,
                volume=62000000,
                yoy_change=20.78,
                yoy_return=13.81,
            ),
        ]

        out_path = write_stock_prices_to_excel(
            ticker="AAPL",
            records=records,
            workbook_path=wb_path,
        )
        self.assertTrue(out_path.is_file())

        # Inspect resulting workbook
        wb_saved = openpyxl.load_workbook(out_path)
        self.assertIn("Stock Price", wb_saved.sheetnames)
        self.assertIn("Income Statement", wb_saved.sheetnames)

        ws_prices = wb_saved["Stock Price"]
        # Header title
        self.assertIn("STOCK PRICES", ws_prices["A1"].value)
        # Column names at row 3
        self.assertEqual(ws_prices.cell(row=3, column=1).value, "Fiscal Year")
        self.assertEqual(
            ws_prices.cell(row=3, column=4).value, "Closing Price ($)"
        )
        # Data rows
        self.assertEqual(ws_prices.cell(row=4, column=1).value, 2022)
        self.assertAlmostEqual(
            ws_prices.cell(row=4, column=4).value, 150.43, places=2
        )
        self.assertEqual(ws_prices.cell(row=5, column=1).value, 2023)
        self.assertAlmostEqual(
            ws_prices.cell(row=5, column=4).value, 171.21, places=2
        )

        # Verification function check
        verified = verify_stock_price_tab(out_path, expected_count=2)
        self.assertTrue(verified)

    def test_run_stock_price_pipeline(self) -> None:
        """Verify full pipeline execution in sandbox mode."""
        wb_path = self.base_path / "AAPL_Test_Analysis.xlsx"
        success = run_stock_price_pipeline(
            ticker="AAPL",
            workbook_path=wb_path,
            data_dir=self.base_path,
            num_years=5,
        )
        self.assertTrue(success)
        self.assertTrue(wb_path.is_file())

        wb = openpyxl.load_workbook(wb_path)
        self.assertIn("Stock Price", wb.sheetnames)


if __name__ == "__main__":
    unittest.main()
